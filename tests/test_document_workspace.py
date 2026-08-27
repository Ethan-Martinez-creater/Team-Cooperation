from __future__ import annotations

import asyncio
import io
import pathlib

import httpx
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool

from coifesp_harness.config import Settings
from coifesp_harness.control_plane import create_app
from coifesp_harness.product import (
    DocumentWorkspaceService,
    NotificationService,
    ProductAccountService,
    ProjectDirectoryService,
    ProjectResourceService,
    TeamCollaborationService,
)
from coifesp_harness.artifacts import (
    ArtifactContentService,
    LocalImmutableArtifactStore,
    SQLAlchemyArtifactRepository,
)

PASSWORD = "Admin-Correct-Horse-42!"


def _make_docx(title: str = "Hello") -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "name"
    sheet["B1"] = 42
    sheet["A2"] = "alpha"
    sheet["B2"] = 7
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _make_pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    textbox = slide.shapes.add_textbox(100000, 100000, 5000000, 500000)
    textbox.text = "Slide one text"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _make_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class _NullAuditLog:
    def __init__(self, engine):
        self.engine = engine

    def append_in_transaction(self, connection, event):
        pass


def stack(tmp_path: pathlib.Path):
    engine = create_engine(
        "sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    accounts = ProductAccountService(engine)
    accounts.create_schema()
    notifications = NotificationService(engine)
    directory = ProjectDirectoryService(engine)
    collaboration = TeamCollaborationService(engine, notifier=notifications)
    artifact_root = tmp_path / "artifacts"
    artifact_root.mkdir()
    artifact_repository = SQLAlchemyArtifactRepository(engine=engine, audit_log=_NullAuditLog(engine))
    artifact_repository.create_schema()
    artifact_content_service = ArtifactContentService(
        artifact_repository,
        LocalImmutableArtifactStore(artifact_root),
    )
    resources = ProjectResourceService(engine, artifact_repository=artifact_repository)
    document_workspace = DocumentWorkspaceService(
        engine=engine,
        resource_service=resources,
        content_service=artifact_content_service,
    )
    app = create_app(
        settings=Settings.from_environment(
            {"COIFESP_ENV": "test", "COIFESP_AUTH_MODE": "builtin"}
        ),
        product_account_service=accounts,
        project_directory_service=directory,
        project_resource_service=resources,
        team_collaboration_service=collaboration,
        notification_service=notifications,
        document_workspace_service=document_workspace,
        artifact_content_service=artifact_content_service,
    )
    team, bootstrap = accounts.register_team(
        team_id="team-doc-alpha", team_handle="doc-alpha", team_name="DocAlpha"
    )
    accounts.change_initial_password(
        login=bootstrap.username,
        current_password=bootstrap.initial_password,
        new_password=PASSWORD,
    )
    session = accounts.login(login=bootstrap.username, password=PASSWORD)
    return app, team.team_id, session.token


async def call(app, method, path, token=None, **kwargs):
    headers = kwargs.pop("headers", {})
    if token:
        headers["Authorization"] = f"Bearer {token}"
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app, raise_app_exceptions=False),
        base_url="http://test",
    ) as client:
        return await client.request(method, path, headers=headers, **kwargs)


def create_project(app, token, name="Doc Project"):
    response = asyncio.run(
        call(
            app,
            "POST",
            "/v1/projects",
            token,
            json={
                "name": name,
                "description": "iteration 3 docs",
                "owner_assignment_name": "产品",
                "owner_kind": "product",
            },
        )
    )
    assert response.status_code == 201, response.text
    return response.json()["project_id"]


def upload_resource(app, token, project_id, filename, content, media_type, title="Sample"):
    response = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources:upload",
            token,
            files={
                "content": (filename, content, media_type),
                "title": (None, title),
                "propagation": (None, "project_readonly"),
            },
            headers={"Idempotency-Key": f"doc-{filename}-{project_id[:8]}"},
        )
    )
    assert response.status_code in (200, 201), response.text
    return response.json()


def test_parse_preview_and_version_chain_for_docx(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    resource = upload_resource(
        app, token, project_id, "sample.docx", _make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/parse", token)
    )
    assert parsed.status_code == 201, parsed.text
    derivative = parsed.json()
    assert derivative["status"] == "ready"

    content = asyncio.run(
        call(
            app,
            "GET",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/derivatives/{derivative['derivative_id']}/content",
            token,
        )
    )
    assert content.status_code == 200, content.text
    body = content.json()
    texts = [item["text"] for item in body["items"]]
    assert any("First paragraph" in text for text in texts)
    assert any("Second paragraph" in text for text in texts)

    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    )
    assert versions.status_code == 200
    assert len(versions.json()) == 1
    assert versions.json()[0]["version_number"] == 1
    assert versions.json()[0]["reason"] == "原始上传"


def test_docx_edit_draft_approve_creates_new_version_original_unchanged(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    original = _make_docx()
    resource = upload_resource(
        app, token, project_id, "edit.docx", original,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/parse", token)
    )
    assert parsed.status_code == 201, parsed.text
    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    ).json()
    source = versions[0]["version_id"]

    draft = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/change-drafts",
            token,
            json={
                "source_version_id": source,
                "modification": {
                    "schema": "coifesp.document-modification.v1",
                    "format": "docx",
                    "operations": [
                        {"op": "replace_paragraph", "paragraph_index": 1, "text": "Updated first."}
                    ],
                },
                "reason": "review fix",
            },
        )
    )
    assert draft.status_code == 201, draft.text
    assert draft.json()["status"] == "pending"

    decided = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/change-drafts/{draft.json()['draft_id']}:decide",
            token,
            json={"approve": True, "expected_version": 1},
        )
    )
    assert decided.status_code == 200, decided.text
    assert decided.json()["status"] == "approved"
    assert decided.json()["generated_version_id"]

    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    ).json()
    assert len(versions) == 2
    assert versions[1]["version_number"] == 2
    assert versions[1]["parent_version_id"] == source

    # The original artifact bytes are unchanged (immutable content addressing).
    from coifesp_harness.artifacts import LocalImmutableArtifactStore
    store = LocalImmutableArtifactStore(tmp_path / "artifacts")
    stored = store.stat(tenant_id=resource["artifact_owner_team_id"], sha256=resource["artifact_sha256"])
    assert stored.size_bytes == len(original)


def test_reject_draft_keeps_version_chain(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    resource = upload_resource(
        app, token, project_id, "reject.docx", _make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/parse", token)
    )
    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    ).json()
    source = versions[0]["version_id"]
    draft = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/change-drafts",
            token,
            json={
                "source_version_id": source,
                "modification": {
                    "schema": "coifesp.document-modification.v1",
                    "format": "docx",
                    "operations": [
                        {"op": "replace_paragraph", "paragraph_index": 0, "text": "Rejected edit."}
                    ],
                },
                "reason": "will reject",
            },
        )
    )
    rejected = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/change-drafts/{draft.json()['draft_id']}:decide",
            token,
            json={"approve": False, "expected_version": 1},
        )
    )
    assert rejected.status_code == 200
    assert rejected.json()["status"] == "rejected"
    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    ).json()
    assert len(versions) == 1


def test_invalid_modification_protocol_rejected(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    resource = upload_resource(
        app, token, project_id, "protocol.docx", _make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/parse", token)
    )
    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    ).json()
    source = versions[0]["version_id"]
    # Unknown operation and unknown top-level field are both rejected.
    for modification in (
        {
            "schema": "coifesp.document-modification.v1",
            "format": "docx",
            "operations": [{"op": "delete_everything"}],
        },
        {
            "schema": "coifesp.document-modification.v1",
            "format": "docx",
            "operations": [{"op": "replace_paragraph", "paragraph_index": 0, "text": "x"}],
            "evil": True,
        },
    ):
        draft = asyncio.run(
            call(
                app,
                "POST",
                f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/change-drafts",
                token,
                json={"source_version_id": source, "modification": modification, "reason": "bad"},
            )
        )
        assert draft.status_code == 422, draft.text


def test_pdf_parse_and_review_only_edit(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    resource = upload_resource(
        app, token, project_id, "note.pdf", _make_pdf(),
        "application/pdf",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/parse", token)
    )
    assert parsed.status_code == 201, parsed.text
    assert parsed.json()["status"] == "ready"

    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    ).json()
    source = versions[0]["version_id"]
    draft = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/change-drafts",
            token,
            json={
                "source_version_id": source,
                "modification": {
                    "schema": "coifesp.document-modification.v1",
                    "format": "pdf",
                    "operations": [{"op": "add_review_note", "page_index": 0, "note": "请复核此页"}],
                },
                "reason": "review",
            },
        )
    )
    assert draft.status_code == 201, draft.text
    decided = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/change-drafts/{draft.json()['draft_id']}:decide",
            token,
            json={"approve": True, "expected_version": 1},
        )
    )
    assert decided.status_code == 200, decided.text
    assert decided.json()["generated_version_id"]
    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions", token)
    ).json()
    assert len(versions) == 2
    assert versions[1]["version_number"] == 2


def test_xlsx_and_pptx_parse_and_edit(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    xlsx = upload_resource(
        app, token, project_id, "data.xlsx", _make_xlsx(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{xlsx['resource_id']}/documents/parse", token)
    )
    assert parsed.status_code == 201, parsed.text
    content = asyncio.run(
        call(
            app,
            "GET",
            f"/v1/projects/{project_id}/resources/{xlsx['resource_id']}/documents/derivatives/{parsed.json()['derivative_id']}/content",
            token,
        )
    ).json()
    texts = [item["text"] for item in content["items"]]
    assert any("alpha" in text for text in texts)

    versions = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{xlsx['resource_id']}/documents/versions", token)
    ).json()
    source = versions[0]["version_id"]
    draft = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{xlsx['resource_id']}/documents/change-drafts",
            token,
            json={
                "source_version_id": source,
                "modification": {
                    "schema": "coifesp.document-modification.v1",
                    "format": "xlsx",
                    "operations": [{"op": "update_cell", "sheet": "Data", "row": 2, "column": 2, "value": 99}],
                },
                "reason": "fix total",
            },
        )
    )
    assert draft.status_code == 201, draft.text
    decided = asyncio.run(
        call(
            app,
            "POST",
            f"/v1/projects/{project_id}/resources/{xlsx['resource_id']}/documents/change-drafts/{draft.json()['draft_id']}:decide",
            token,
            json={"approve": True, "expected_version": 1},
        )
    )
    assert decided.status_code == 200, decided.text

    pptx = upload_resource(
        app, token, project_id, "deck.pptx", _make_pptx(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{pptx['resource_id']}/documents/parse", token)
    )
    assert parsed.status_code == 201, parsed.text
    content = asyncio.run(
        call(
            app,
            "GET",
            f"/v1/projects/{project_id}/resources/{pptx['resource_id']}/documents/derivatives/{parsed.json()['derivative_id']}/content",
            token,
        )
    ).json()
    texts = [item["text"] for item in content["items"]]
    assert any("Slide one text" in text for text in texts)


def test_parse_failure_degrades_gracefully(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    # A corrupt DOCX (random bytes) must not leak an internal stack trace.
    resource = upload_resource(
        app, token, project_id, "broken.docx", b"\x00\x01not a zip",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    parsed = asyncio.run(
        call(app, "POST", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/parse", token)
    )
    assert parsed.status_code == 201, parsed.text
    body = parsed.json()
    assert body["status"] == "failed"
    assert body["error_category"]  # stable category, no stack
    assert "Traceback" not in parsed.text

    # The original file is still downloadable.
    download = asyncio.run(
        call(app, "GET", f"/v1/projects/{project_id}/resources/{resource['resource_id']}/content", token)
    )
    assert download.status_code == 200
    assert download.content == b"\x00\x01not a zip"


def test_cross_team_cannot_touch_document(tmp_path):
    app, team_id, token = stack(tmp_path)
    project_id = create_project(app, token)
    resource = upload_resource(
        app, token, project_id, "private.docx", _make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    other = asyncio.run(
        call(app, "POST", "/v1/teams/register", json={"handle": "beta-doc", "name": "BetaDoc"})
    )
    other_details = other.json()
    asyncio.run(
        call(
            app,
            "POST",
            "/v1/accounts/change-initial-password",
            json={
                "login": other_details["administrator_username"],
                "current_password": other_details["administrator_initial_password"],
                "new_password": PASSWORD,
            },
        )
    )
    other_session = asyncio.run(
        call(
            app,
            "POST",
            "/v1/sessions",
            json={"login": other_details["administrator_username"], "password": PASSWORD},
        )
    )
    other_token = other_session.json()["access_token"]
    versions = asyncio.run(
        call(
            app,
            "GET",
            f"/v1/projects/{project_id}/resources/{resource['resource_id']}/documents/versions",
            other_token,
        )
    )
    assert versions.status_code in (403, 404)
